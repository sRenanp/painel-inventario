"""Geração da planilha Excel."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from analise import fmt_config, fmt_gb, fmt_upgrade, resumo_configuracoes


def _seguro(v):
    # Evita que nomes vindos do OCS (ex.: "=HYPERLINK(...)") sejam interpretados como fórmula
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def _aba(wb, titulo, cabecalho, linhas):
    ws = wb.create_sheet(titulo)
    ws.append(cabecalho)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")
    for linha in linhas:
        ws.append([_seguro(v) for v in linha])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, start=1):
        largura = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 70)


def gerar(analises, grupos, itens_estoque, alertas) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    _aba(wb, "Máquinas", ["Máquina", "Modelo", "Service tag", "Usuário", "Pentes", "Qtd pentes", "Total (GB)",
                          "Slots", "Slots livres", "Máx. suportado (GB)", "Tipo", "Upgrades possíveis",
                          "Último inventário"], [
        [a["maquina"].nome, a["maquina"].modelo, a["maquina"].service_tag, a["maquina"].usuario, a["config"],
         len(a["pentes_gb"]), a["total_gb"], a["slots"], a["livres"], a["max_gb"], a["tipo"],
         "\n".join(fmt_upgrade(u) for u in a["upgrades"]), a["maquina"].ultimo_inventario]
        for a in analises])
    for c in wb["Máquinas"]["M"][1:]:  # Último inventário
        c.number_format = "dd/mm/yyyy hh:mm"

    _aba(wb, "Por modelo", ["Modelo", "Qtd", "Média (GB)", "Configurações", "Com slot livre", "Com upgrade",
                            "Máx. suportado"], [
        [g["modelo"], g["qtd"], round(g["media_gb"], 1),
         "; ".join(f"{cfg}: {n}" for cfg, n in g["configs"]), g["com_slot_livre"], g["com_upgrade"],
         fmt_gb(g["spec"].get("max_gb")) if g["spec"] else "sem spec"]
        for g in grupos])

    _aba(wb, "Configurações", ["Modelo", "Total (GB)", "Pentes", "Slots livres", "Máquinas"], [
        [c["modelo"], c["total_gb"], c["config"], c["livres"], c["qtd"]] for c in resumo_configuracoes(analises)])

    _aba(wb, "Estoque", ["Service tag", "Modelo", "Situação", "Local", "Observação", "Atualizado em"], [
        [i["service_tag"], i["modelo"], i["situacao"], i["local"], i["observacao"], i["atualizado_em"]]
        for i in itens_estoque])

    _aba(wb, "Divergências", ["Gravidade", "Tipo", "Service tag", "Máquina", "Detalhe"], [
        [d["gravidade"], d["tipo"], d["service_tag"], d["maquina"], d["detalhe"]] for d in alertas])

    return _salvar(wb)


def gerar_compras(pentes, plano, meta_gb) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    _aba(wb, f"Comprar (meta {meta_gb}GB)", ["Tipo", "Tamanho", "Necessários", "Reaproveitados das trocas",
                                             "Comprar", "Sobram"], [
        [l["tipo"], fmt_gb(l["gb"]), l["necessarios"], min(l["necessarios"], l["liberados"]), l["comprar"],
         l["sobram"]] for l in plano["resumo"]])

    _aba(wb, "Por máquina", ["Máquina", "Modelo", "Service tag", "Usuário", "Atual", "Alvo", "Comprar",
                             "Pentes retirados", "Obs"], [
        [p["analise"]["maquina"].nome, p["analise"]["maquina"].modelo, p["analise"]["maquina"].service_tag,
         p["analise"]["maquina"].usuario, p["analise"]["config"], fmt_gb(p["alvo"]), fmt_config(p["comprar"]),
         fmt_config(p["liberados"]) if p["liberados"] else "", p["obs"]] for p in plano["maquinas"]])

    _aba(wb, "Pentes instalados", ["Tipo", "Tamanho", "Quantidade"], [
        [tipo, fmt_gb(gb), n] for (tipo, gb), n in pentes["por_tipo"]])

    _aba(wb, "Pentes por modelo", ["Modelo"] + [fmt_gb(t) for t in pentes["tamanhos"]] + ["Total"], [
        [modelo] + [c.get(t, 0) for t in pentes["tamanhos"]] + [sum(c.values())]
        for modelo, c in pentes["por_modelo"]])
    return _salvar(wb)


def _salvar(wb) -> io.BytesIO:
    saida = io.BytesIO()
    wb.save(saida)
    saida.seek(0)
    return saida
